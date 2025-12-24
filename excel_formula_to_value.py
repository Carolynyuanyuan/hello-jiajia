#!/usr/bin/env python3
"""
Excel公式转数值工具
将Excel文件中的所有公式转换为计算后的数值，相当于"复制-仅粘贴数值"操作
支持保留原始格式、样式和工作表结构
"""

import openpyxl
from openpyxl.utils import get_column_letter
import argparse
import sys
from pathlib import Path


def convert_formulas_to_values(input_file, output_file=None):
    """
    将Excel文件中的公式转换为数值

    参数:
        input_file: 输入Excel文件路径
        output_file: 输出Excel文件路径（如果为None，则自动生成）

    返回:
        output_file: 输出文件路径
    """
    try:
        # 加载工作簿，data_only=False以获取公式
        print(f"正在加载文件: {input_file}")
        wb = openpyxl.load_workbook(input_file, data_only=False)

        # 同时加载一个data_only=True的版本来获取计算值
        wb_values = openpyxl.load_workbook(input_file, data_only=True)

        total_formulas = 0

        # 遍历所有工作表
        for sheet_name in wb.sheetnames:
            print(f"处理工作表: {sheet_name}")
            ws = wb[sheet_name]
            ws_values = wb_values[sheet_name]
            sheet_formula_count = 0

            # 遍历所有单元格
            for row in ws.iter_rows():
                for cell in row:
                    # 检查单元格是否包含公式
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                        # 获取对应的计算值
                        value_cell = ws_values[cell.coordinate]
                        calculated_value = value_cell.value

                        # 替换公式为计算值
                        cell.value = calculated_value
                        sheet_formula_count += 1

            print(f"  转换了 {sheet_formula_count} 个公式单元格")
            total_formulas += sheet_formula_count

        # 生成输出文件名
        if output_file is None:
            input_path = Path(input_file)
            output_file = input_path.parent / f"{input_path.stem}_values{input_path.suffix}"

        # 保存新文件
        print(f"\n正在保存到: {output_file}")
        wb.save(output_file)

        print(f"\n完成！总共转换了 {total_formulas} 个公式单元格")
        print(f"输出文件: {output_file}")

        return output_file

    except FileNotFoundError:
        print(f"错误: 找不到文件 '{input_file}'", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"错误: {str(e)}", file=sys.stderr)
        sys.exit(1)


def main():
    parser = argparse.ArgumentParser(
        description='将Excel文件中的公式转换为数值',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
使用示例:
  # 转换单个文件（自动生成输出文件名）
  python excel_formula_to_value.py input.xlsx

  # 指定输出文件名
  python excel_formula_to_value.py input.xlsx -o output.xlsx

  # 使用简短参数
  python excel_formula_to_value.py data.xlsx -o data_clean.xlsx
        """
    )

    parser.add_argument(
        'input_file',
        help='输入Excel文件路径'
    )

    parser.add_argument(
        '-o', '--output',
        dest='output_file',
        help='输出Excel文件路径（可选，默认为输入文件名_values.xlsx）'
    )

    args = parser.parse_args()

    convert_formulas_to_values(args.input_file, args.output_file)


if __name__ == "__main__":
    main()
