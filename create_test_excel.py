#!/usr/bin/env python3
"""
创建一个包含公式的测试Excel文件
用于演示excel_formula_to_value.py的功能
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from datetime import datetime


def create_test_excel(filename='test_formulas.xlsx'):
    """创建一个包含各种公式的测试Excel文件"""

    wb = openpyxl.Workbook()

    # ========== 工作表1: 销售数据 ==========
    ws1 = wb.active
    ws1.title = "销售数据"

    # 设置标题样式
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 添加标题行
    headers = ['产品', '单价', '数量', '小计', '税率', '含税金额']
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    # 添加数据和公式
    products = [
        ('苹果', 10, 50),
        ('香蕉', 8, 30),
        ('橙子', 12, 40),
        ('葡萄', 15, 25),
        ('西瓜', 20, 15),
    ]

    for row, (product, price, qty) in enumerate(products, 2):
        ws1.cell(row=row, column=1, value=product)
        ws1.cell(row=row, column=2, value=price)
        ws1.cell(row=row, column=3, value=qty)

        # 小计 = 单价 × 数量（公式）
        ws1.cell(row=row, column=4, value=f'=B{row}*C{row}')

        # 税率
        ws1.cell(row=row, column=5, value=0.13)

        # 含税金额 = 小计 × (1 + 税率)（公式）
        ws1.cell(row=row, column=6, value=f'=D{row}*(1+E{row})')

    # 添加汇总行
    summary_row = len(products) + 2
    ws1.cell(row=summary_row, column=1, value='总计').font = Font(bold=True)
    ws1.cell(row=summary_row, column=4, value=f'=SUM(D2:D{summary_row-1})')
    ws1.cell(row=summary_row, column=6, value=f'=SUM(F2:F{summary_row-1})')

    # 调整列宽
    ws1.column_dimensions['A'].width = 12
    ws1.column_dimensions['B'].width = 10
    ws1.column_dimensions['C'].width = 10
    ws1.column_dimensions['D'].width = 12
    ws1.column_dimensions['E'].width = 10
    ws1.column_dimensions['F'].width = 12

    # ========== 工作表2: 统计分析 ==========
    ws2 = wb.create_sheet("统计分析")

    # 标题
    ws2['A1'] = '统计指标'
    ws2['B1'] = '值'
    ws2['A1'].font = Font(bold=True, size=12)
    ws2['B1'].font = Font(bold=True, size=12)

    # 各种统计公式
    stats = [
        ('总销售额', '=销售数据!D7'),
        ('平均单价', '=AVERAGE(销售数据!B2:B6)'),
        ('最高单价', '=MAX(销售数据!B2:B6)'),
        ('最低单价', '=MIN(销售数据!B2:B6)'),
        ('总数量', '=SUM(销售数据!C2:C6)'),
        ('平均数量', '=AVERAGE(销售数据!C2:C6)'),
        ('数据条数', '=COUNTA(销售数据!A2:A6)'),
    ]

    for row, (label, formula) in enumerate(stats, 2):
        ws2.cell(row=row, column=1, value=label)
        ws2.cell(row=row, column=2, value=formula)

    ws2.column_dimensions['A'].width = 15
    ws2.column_dimensions['B'].width = 15

    # ========== 工作表3: 日期计算 ==========
    ws3 = wb.create_sheet("日期计算")

    ws3['A1'] = '项目'
    ws3['B1'] = '开始日期'
    ws3['C1'] = '结束日期'
    ws3['D1'] = '持续天数'
    ws3['E1'] = '状态'

    # 设置标题格式
    for col in ['A', 'B', 'C', 'D', 'E']:
        ws3[f'{col}1'].font = Font(bold=True)
        ws3[f'{col}1'].fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")

    # 添加项目数据
    projects = [
        ('项目A', '2024-01-01', '2024-03-31'),
        ('项目B', '2024-02-15', '2024-06-30'),
        ('项目C', '2024-04-01', '2024-08-31'),
    ]

    for row, (project, start, end) in enumerate(projects, 2):
        ws3.cell(row=row, column=1, value=project)
        ws3.cell(row=row, column=2, value=start)
        ws3.cell(row=row, column=3, value=end)

        # 计算天数（公式）
        ws3.cell(row=row, column=4, value=f'=C{row}-B{row}')

        # 状态判断（公式：如果结束日期>今天，显示"进行中"，否则"已完成"）
        ws3.cell(row=row, column=5, value=f'=IF(C{row}>TODAY(),"进行中","已完成")')

    ws3.column_dimensions['A'].width = 12
    ws3.column_dimensions['B'].width = 15
    ws3.column_dimensions['C'].width = 15
    ws3.column_dimensions['D'].width = 12
    ws3.column_dimensions['E'].width = 12

    # 保存文件
    wb.save(filename)
    print(f"✓ 测试文件已创建: {filename}")
    print(f"\n文件包含以下工作表:")
    print(f"  1. 销售数据 - 包含乘法、加法、求和公式")
    print(f"  2. 统计分析 - 包含引用、平均值、最大/最小值公式")
    print(f"  3. 日期计算 - 包含日期运算和条件判断公式")
    print(f"\n你可以运行以下命令来测试转换:")
    print(f"  python excel_formula_to_value.py {filename}")


if __name__ == "__main__":
    create_test_excel()
